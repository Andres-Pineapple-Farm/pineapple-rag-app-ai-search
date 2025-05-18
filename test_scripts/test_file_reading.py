"""
Test script to diagnose file opening behavior with binary files
"""
import os
from pathlib import Path
import logging

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

def test_file_reading(file_path):
    """Test reading a file in different ways"""
    logger.info(f"Testing file reading for: {file_path}")
    
    # Check if file exists
    if not os.path.exists(file_path):
        logger.error(f"File not found: {file_path}")
        return
        
    # Get file size
    file_size = os.path.getsize(file_path)
    logger.info(f"File size: {file_size} bytes")
    
    # Try to read as text
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
            logger.info(f"Successfully read as text with UTF-8 encoding: {len(content)} chars")
    except UnicodeDecodeError:
        logger.warning("Failed to read as text with UTF-8 encoding - this is expected for binary files")
    except Exception as e:
        logger.error(f"Error reading as text: {type(e).__name__}: {e}")
    
    # Try to read as binary
    try:
        with open(file_path, 'rb') as f:
            binary_content = f.read()
            logger.info(f"Successfully read as binary: {len(binary_content)} bytes")
    except Exception as e:
        logger.error(f"Error reading as binary: {type(e).__name__}: {e}")

    # Try specialized handling based on file type
    file_ext = Path(file_path).suffix.lower()
    if file_ext in ['.docx', '.doc']:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs_count = len(doc.paragraphs)
            logger.info(f"Successfully read as Word document: {paragraphs_count} paragraphs")
        except ImportError:
            logger.warning("python-docx not installed")
        except Exception as e:
            logger.error(f"Error reading as Word document: {type(e).__name__}: {e}")
    
    elif file_ext in ['.pptx', '.ppt']:
        try:
            import pptx
            presentation = pptx.Presentation(file_path)
            slides_count = len(presentation.slides)
            logger.info(f"Successfully read as PowerPoint presentation: {slides_count} slides")
        except ImportError:
            logger.warning("python-pptx not installed")
        except Exception as e:
            logger.error(f"Error reading as PowerPoint presentation: {type(e).__name__}: {e}")

if __name__ == "__main__":
    # Test with sample files
    sample_docx = "sample_docs/Commercial_Office_Lease_Agreement.docx"
    sample_pptx = "sample_docs/Industrial_Real_Estate_Trends_Midwest_2020_2025.pptx"
    
    test_file_reading(sample_docx)
    test_file_reading(sample_pptx)
